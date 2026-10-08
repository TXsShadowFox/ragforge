"""Make small test files in memory: PDFs with text, and DOCX files."""

import io
from collections.abc import Sequence

import docx


def make_pdf(pages: Sequence[str]) -> bytes:
    """A small valid PDF: one page per string, one text line per line of the string.

    Written by hand (it is a simple format), because pypdf cannot add text to new pages.
    Uses Helvetica, a font every PDF reader has, so only latin-1 characters work. The
    font says WinAnsiEncoding: without it, "'" is read back as a curly quote (or garbage).
    """
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    page_ids: list[int] = []
    for number, text in enumerate(pages):
        content_id, page_id = 4 + 2 * number, 5 + 2 * number
        content = _text_operations(text)
        objects[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content)
        objects[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id
        )
        page_ids.append(page_id)
    kids = b" ".join(b"%d 0 R" % page_id for page_id in page_ids)
    objects[2] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))
    return _write_pdf(objects)


def handbook_page(page: int) -> str:
    """About 110 words. Only its middle line is special: it names `zone<page>`.

    So a question about "zone37" has exactly one right page: 37.
    """
    filler = "Students should read this handbook carefully and follow every rule on campus."
    special = f"Only students with a blue pass may enter zone{page} after dark."
    return "\n".join([filler] * 4 + [special] + [filler] * 4)


def make_docx(paragraphs: Sequence[str], table: Sequence[Sequence[str]] = ()) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(word_table.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _text_operations(text: str) -> bytes:
    lines = [b"BT", b"/F1 10 Tf", b"14 TL", b"50 750 Td"]
    lines += [b"(%s) Tj T*" % _pdf_string(line) for line in text.split("\n")]
    lines.append(b"ET")
    return b"\n".join(lines)


def _pdf_string(text: str) -> bytes:
    raw = text.encode("latin-1", "replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def _write_pdf(objects: dict[int, bytes]) -> bytes:
    """Objects, then the cross-reference table (where each object starts), then the trailer."""
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(out)
        out += b"%d 0 obj\n%s\nendobj\n" % (object_id, objects[object_id])
    xref_start = len(out)
    size = max(objects) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    out += b"".join(b"%010d 00000 n \n" % offsets[object_id] for object_id in range(1, size))
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref_start)
    return bytes(out)
